# Contenido único -> HTML (PDF A4 vía Chromium) y DOCX
import html, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT='/home/user/Molu/master-sistemas-inteligentes/metodologia-de-la-investigacion/practica-del-problema-al-articulo/paso-1-planteamiento/'
FS=float(sys.argv[1]) if len(sys.argv)>1 else 11
N=' '
TITLE='ESTIMACIÓN DE LA CALIDAD DE TRANSMISIÓN MEDIANTE APRENDIZAJE AUTOMÁTICO EN LA RED DWDM FERROVIARIA MOSCÚ-KAZÁNSKAYA – RIAZÁN'
META=['Angel Luis Acosta González','Máster Universitario en Sistemas Inteligentes (USAL) · Metodología de la Investigación','Paso 1: Planteamiento del problema, pregunta e hipótesis · 28 de septiembre de 2026']
# bloques: ('h',texto) | ('p',[(txt,fmt)]) | ('li',[...]) ; fmt: ''|'b'|'i'|'sup'
B=[
('h','1. PLANTEAMIENTO DEL PROBLEMA'),
('p',[('Este trabajo parte del diseño que propuse en mi trabajo de fin de grado (RUT MIIT, 2026) para la red primaria de comunicaciones del tramo Moscú-Kazánskaya – Riazán-1 (198,3 km): un anillo DWDM entre Moscú-Kazánskaya, Voskresensk y Riazán-1, más una cadena de acceso de 22 estaciones, con canales de 10 Gbit/s sobre fibra G.652 (α'+N+'='+N+'0,22 dB/km). La calidad de transmisión (QoT) se verificó con un modelo analítico aplicado solo al tramo más desfavorable (Voskresensk – Riazán-1, 108,9 km). El resultado fue Q'+N+'='+N+'7,37 y BER'+N+'='+N+'8,9·10',''),('−14','sup'),(', frente a un umbral normativo de 10',''),('−11','sup'),(' (Q'+N+'='+N+'6,71).','')]),
('p',[('El enfoque utilizado en ese proyecto es estático y conservador: se calcula una sola vez, para un caso extremo y con márgenes fijos. No tiene en cuenta la degradación durante la explotación: envejecimiento de la fibra y de los amplificadores EDFA, empalmes por reparaciones o variaciones de temperatura. Tampoco recoge el hielo y las inundaciones, que el propio proyecto identifica como riesgos del trazado, ni la reconfiguración de canales en los OADM. Recalcular a mano cada canal óptico (',''),('lightpath','i'),(') en cada situación es costoso, y los márgenes fijos tienden a sobredimensionar la red.','')]),
('h','2. PREGUNTA INICIAL DE INVESTIGACIÓN'),
('p',[('¿Puede un modelo de aprendizaje automático, entrenado con parámetros que el operador puede medir, predecir si cada canal de la red ferroviaria cumple el umbral de calidad BER'+N+'≤'+N+'10',''),('−11','sup'),(' en condiciones de degradación, con más acierto que un modelo sencillo?','')]),
('h','3. NOVEDAD E INTERÉS'),
('li',[('Vacío: ','b'),('la estimación de QoT con aprendizaje automático está consolidada en redes troncales de larga distancia [1, 2], incluso con modelos entrenados con datos sintéticos [3]. Sin embargo, no se han encontrado trabajos que la apliquen a redes de comunicaciones ferroviarias, lo que se confirmará en la revisión bibliográfica del Paso 2. Estas redes son regionales, con muchas estaciones de acceso, expuestas a condiciones ambientales severas y, en este caso, limitadas a equipos de fabricación nacional.','')]),
('li',[('Interés: ','b'),('se sustituye un cálculo puntual por un sistema inteligente que evalúa la red de forma continua. Es un paso hacia la gestión autónoma de redes y el mantenimiento predictivo de una infraestructura crítica para la seguridad ferroviaria.','')]),
('h','4. OBJETIVOS'),
('p',[('General: ','b'),('desarrollar y evaluar un estimador de QoT basado en aprendizaje automático para la red DWDM del tramo Moscú-Kazánskaya – Riazán.','')]),
('p',[('1. Construir un conjunto de datos de canales ópticos de la red con un modelo físico de transmisión basado en el del proyecto, incluyendo condiciones de degradación.','')]),
('p',[('2. Entrenar y comparar distintos modelos de aprendizaje automático con un modelo de referencia sencillo.','')]),
('p',[('3. Evaluar su acierto en canales y condiciones no vistos durante el entrenamiento.','')]),
('p',[('4. Cuantificar el margen de diseño que podría ahorrarse sin incumplir el umbral de BER.','')]),
('h','5. HIPÓTESIS PROVISIONAL'),
('p',[('Si se entrena un modelo de ensamble (por ejemplo, ',''),('Random Forest','i'),(') con parámetros que el operador puede medir, será posible predecir si un canal cumple el umbral BER'+N+'≤'+N+'10',''),('−11','sup'),(' con mayor acierto que un modelo lineal sencillo.','')]),
('p',[('Limitación prevista: ','b'),('los datos no serán mediciones reales de la red, sino que se generarán por ordenador con ecuaciones físicas de transmisión óptica (datos sintéticos). Por ello, los resultados deberán contrastarse con conjuntos de datos públicos medidos en redes reales.','')]),
]
REFS=[
[('[1] Pointurier, Y. (2021). Machine learning techniques for quality of transmission estimation in optical networks. ',''),('Journal of Optical Communications and Networking','i'),(', 13(4), B60. https://doi.org/10.1364/JOCN.417434','')],
[('[2] Kozdrowski, S., Cichosz, P., Paziewski, P., & Sujecki, S. (2021). Machine learning algorithms for prediction of the quality of transmission in optical networks. ',''),('Entropy','i'),(', 23(1), 7. https://doi.org/10.3390/e23010007','')],
[('[3] Rottondi, C., Barletta, L., Giusti, A., & Tornatore, M. (2018). Machine-learning method for quality of transmission prediction of unestablished lightpaths. ',''),('Journal of Optical Communications and Networking','i'),(', 10(2), A286. https://doi.org/10.1364/JOCN.10.00A286','')],
]
def hr(runs):
    o=''
    for t,f in runs:
        t=html.escape(t)
        o+= {'b':f'<b>{t}</b>','i':f'<i>{t}</i>','sup':f'<sup>{t}</sup>'}.get(f,t)
    return o
body=f'<p class="t">{TITLE}</p>'+''.join(f'<p class="m">{html.escape(m)}</p>' for m in META)+'<div class="sep"></div>'
for k,v in B:
    if k=='h': body+=f'<p class="h">{v}</p>'
    elif k=='p': body+=f'<p class="p">{hr(v)}</p>'
    else: body+=f'<p class="p">•&nbsp;{hr(v)}</p>'
body+='<p class="h">REFERENCIAS</p>'+''.join(f'<p class="r">{hr(r)}</p>' for r in REFS)
css=f'''@page{{size:A4;margin:1.7cm 2cm 1.5cm 2cm}}
body{{font-family:"Times New Roman",Times,serif;font-size:{FS}pt;line-height:1.1;margin:0;color:#000}}
p{{margin:0}} .t{{text-align:center;font-weight:bold;margin-bottom:4pt}} .m{{text-align:center;font-size:{FS-1}pt}}
.sep{{border-bottom:.8pt solid #000;margin:5pt 0 2pt}}
.h{{font-weight:bold;margin-top:7pt;margin-bottom:2pt}} .p{{text-align:justify;text-indent:1.25cm;hyphens:auto}}
.r{{font-size:{FS-2}pt;padding-left:1em;text-indent:-1em;margin-bottom:1pt}} sup{{line-height:0}}'''
open(OUT+'ficha-paso-1.html','w').write(f'<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Ficha Paso 1</title><style>{css}</style></head><body>{body}</body></html>')

d=Document(); s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.7); s.bottom_margin=Cm(1.5); s.left_margin=s.right_margin=Cm(2)
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(FS); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
st.paragraph_format.space_after=Pt(0); st.paragraph_format.space_before=Pt(0); st.paragraph_format.line_spacing=1.1
def add(runs,align=WD_ALIGN_PARAGRAPH.JUSTIFY,indent=True,size=None,bold=False):
    p=d.add_paragraph(); p.alignment=align
    if indent: p.paragraph_format.first_line_indent=Cm(1.25)
    for t,f in runs:
        r=p.add_run(t); r.bold=bold or f=='b'; r.italic=f=='i'; r.font.superscript=f=='sup'
        if size: r.font.size=Pt(size)
    return p
p=add([(TITLE,'')],WD_ALIGN_PARAGRAPH.CENTER,False,bold=True); p.paragraph_format.space_after=Pt(4)
for m in META: add([(m,'')],WD_ALIGN_PARAGRAPH.CENTER,False,FS-1)
from docx.oxml import OxmlElement
pPr=d.paragraphs[-1]._p.get_or_add_pPr(); bd=OxmlElement('w:pBdr'); e=OxmlElement('w:bottom')
for a,v in [('val','single'),('sz','6'),('space','4'),('color','000000')]: e.set(qn('w:'+a),v)
bd.append(e); sp=pPr.find(qn('w:spacing')); (sp.addprevious(bd) if sp is not None else pPr.insert(0,bd)) if pPr.find(qn('w:jc')) is None else pPr.find(qn('w:jc')).addprevious(bd)
for k,v in B:
    if k=='h':
        p=add([(v,'')],WD_ALIGN_PARAGRAPH.LEFT,False,bold=True); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
    elif k=='p': add(v)
    else: add([('• ','')]+v)
p=add([('REFERENCIAS','')],WD_ALIGN_PARAGRAPH.LEFT,False,bold=True); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2)
for r in REFS:
    p=add(r,WD_ALIGN_PARAGRAPH.LEFT,False,FS-2); p.paragraph_format.left_indent=Pt(FS); p.paragraph_format.first_line_indent=Pt(-FS)
z=d.settings.element.find(qn('w:zoom'))
if z is not None: z.set(qn('w:percent'),'100')
order=['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','spacing','ind','contextualSpacing','jc','rPr']
rank={qn('w:'+n):i for i,n in enumerate(order)}
for pp in d.element.body.iter(qn('w:pPr')):
    k=list(pp)
    for e in k: pp.remove(e)
    for e in sorted(k,key=lambda e:rank.get(e.tag,99)): pp.append(e)
d.save(OUT+'ficha-paso-1.docx')
