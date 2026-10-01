"""Genera borrador-articulo.pdf y borrador-articulo.docx a partir de contenido.py.
Uso: python generar_borrador.py   (el PDF necesita Chrome/Chromium; ruta en la variable CHROME)"""
import os, re, subprocess
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from contenido import BLOQUES

AQUI = Path(__file__).parent
CHROME = os.environ.get("CHROME", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
TITULO = next(t for k, t, *_ in BLOQUES if k == "titulo")

# ---------------------------------------------------------------- PDF (HTML + Chromium)
def html():
    out = []
    for b in BLOQUES:
        k = b[0]
        if k == "titulo": out.append(f"<h1>{b[1]}</h1>")
        elif k == "autor": out.append(f'<div class="autor">{b[1]}</div>')
        elif k == "nota": out.append(f'<div class="nota">{b[1]}</div>')
        elif k == "resumen": out.append(f'<div class="abs"><b>Resumen.</b> {b[1]}</div>')
        elif k == "claves": out.append(f'<div class="abs"><b>Palabras clave:</b> {b[1]}</div>')
        elif k == "h2": out.append(f"<h2>{b[1]}</h2>")
        elif k == "h3": out.append(f"<h3>{b[1]}</h3>")
        elif k == "p": out.append(f"<p>{b[1]}</p>")
        elif k == "eq": out.append(f'<div class="eq"><span>{b[1]}</span><span>({b[2]})</span></div>')
        elif k == "tabla":
            filas = b[2]
            cab = "".join(f"<th>{c}</th>" for c in filas[0])
            cuerpo = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in f) + "</tr>" for f in filas[1:])
            out.append(f'<div class="cap">{b[1]}</div><table><thead><tr>{cab}</tr></thead><tbody>{cuerpo}</tbody></table>')
        elif k == "fig":
            out.append(f'<figure><img src="{b[1]}"><figcaption>{b[2]}</figcaption></figure>')
        elif k == "refs":
            out.append("<h2>Referencias</h2><ol class='refs'>" + "".join(f"<li>{r}.</li>" for r in b[1]) + "</ol>")
    css = """@page{size:A4;margin:2.2cm 2.4cm}
body{font-family:"Times New Roman",Times,serif;font-size:10.5pt;line-height:1.3;margin:0;color:#000}
h1{font-size:14pt;text-align:center;margin:0 0 6pt} .autor{text-align:center;font-size:10pt;margin-bottom:8pt}
.nota{font-size:8.5pt;font-style:italic;text-align:center;color:#333;margin-bottom:10pt}
.abs{font-size:9pt;text-align:justify;margin:0 1.2cm 5pt}
h2{font-size:11.5pt;margin:12pt 0 4pt;break-after:avoid} h3{font-size:10.5pt;margin:8pt 0 3pt;break-after:avoid}
p{text-align:justify;margin:0 0 4pt;text-indent:1em}
.eq{display:flex;justify-content:space-between;margin:4pt 0 6pt 2.5cm;break-inside:avoid}
.cap{font-size:9pt;margin:8pt 0 2pt;text-align:center;break-after:avoid}
table{border-collapse:collapse;margin:0 auto 8pt;font-size:8.5pt;break-inside:avoid}
th,td{padding:1.5pt 5pt;text-align:center} td:nth-child(-n+2),th:nth-child(-n+2){text-align:left}
thead tr{border-top:1pt solid #000;border-bottom:.6pt solid #000} tbody tr:last-child{border-bottom:1pt solid #000}
figure{margin:6pt 0 8pt;text-align:center;break-inside:avoid} figure img{width:46%}
figcaption{font-size:9pt;margin-top:2pt} .refs{font-size:9pt;padding-left:1.6em} .refs li{margin-bottom:2pt}"""
    return f'<!doctype html><html lang="es"><meta charset="utf-8"><title>{TITULO}</title><style>{css}</style><body>{"".join(out)}</body></html>'

def pdf():
    tmp = AQUI / "borrador-articulo.html"
    tmp.write_text(html(), encoding="utf-8")
    subprocess.run([CHROME, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={AQUI / 'borrador-articulo.pdf'}", tmp.as_uri()], check=True, capture_output=True)
    tmp.unlink()
    try:
        from pypdf import PdfWriter
        w = PdfWriter(clone_from=AQUI / "borrador-articulo.pdf")
        w.add_metadata({"/Title": TITULO, "/Author": "Angel L. Acosta-González"})
        w.write(AQUI / "borrador-articulo.pdf")
    except ImportError:
        pass

# ---------------------------------------------------------------- DOCX
class Runs(HTMLParser):
    """Convierte texto con <b>, <i>, <sup>, <sub>, <br> en fragmentos con formato."""
    def __init__(self):
        super().__init__(); self.estado = set(); self.trozos = []
    def handle_starttag(self, tag, attrs):
        if tag == "br": self.trozos.append(("\n", frozenset()))
        else: self.estado.add(tag)
    def handle_endtag(self, tag): self.estado.discard(tag)
    def handle_data(self, data): self.trozos.append((data, frozenset(self.estado)))

def escribir(par, texto, size=None, bold=False, italic=False):
    r = Runs(); r.feed(texto)
    for t, fmt in r.trozos:
        run = par.add_run(t)
        run.bold = bold or "b" in fmt
        run.italic = italic or "i" in fmt
        run.font.superscript = "sup" in fmt
        run.font.subscript = "sub" in fmt
        if size: run.font.size = Pt(size)

def docx():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10.5)
    for s in d.sections: s.left_margin = s.right_margin = Cm(2.4); s.top_margin = s.bottom_margin = Cm(2.2)
    for b in BLOQUES:
        k = b[0]
        if k in ("titulo", "autor", "nota"):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            escribir(p, b[1], size={"titulo": 14, "autor": 10, "nota": 8.5}[k], bold=k == "titulo", italic=k == "nota")
        elif k in ("resumen", "claves"):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(1.2)
            escribir(p, ("<b>Resumen.</b> " if k == "resumen" else "<b>Palabras clave:</b> ") + b[1], size=9)
        elif k in ("h2", "h3"):
            p = d.add_paragraph(); escribir(p, b[1], size=11.5 if k == "h2" else 10.5, bold=True)
        elif k == "p":
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0.4); escribir(p, b[1])
        elif k == "eq":
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; escribir(p, f"{b[1]}        ({b[2]})")
        elif k == "tabla":
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; escribir(p, b[1], size=9)
            filas = b[2]; t = d.add_table(rows=len(filas), cols=len(filas[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.style = "Table Grid"
            for i, f in enumerate(filas):
                for j, c in enumerate(f):
                    celda = t.cell(i, j); celda.text = ""; escribir(celda.paragraphs[0], c, size=8.5, bold=i == 0)
        elif k == "fig":
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(AQUI / b[1]), width=Cm(7.5))
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; escribir(p, b[2], size=9)
        elif k == "refs":
            p = d.add_paragraph(); escribir(p, "Referencias", size=11.5, bold=True)
            for i, r in enumerate(b[1], 1):
                p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.6); escribir(p, f"{i}. {r}.", size=9)
    d.core_properties.title = re.sub("<[^>]+>", "", TITULO)
    d.core_properties.author = "Angel L. Acosta-González"
    d.save(AQUI / "borrador-articulo.docx")

if __name__ == "__main__":
    pdf(); docx()
    print("borrador-articulo.pdf y borrador-articulo.docx creados")
